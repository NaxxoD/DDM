# tools/scripts/ddm_code_indexer.py
from __future__ import annotations

import argparse
import ast
import json
import re
import sys
from dataclasses import dataclass, asdict
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

# ---------------------------
# Defaults / config
# ---------------------------

DEFAULT_EXCLUDES = {
    ".git", "__pycache__", ".venv", "venv", "env",
    "logs", "log_data", "autorun_log_data", "reports", "legacy",
    ".mypy_cache", ".pytest_cache",
}

DEFAULT_FORMATS = ["md", "json"]  # interactive default
DEFAULT_OUT_DIR = "docs"          # relative to root

PY_EXT = ".py"

# Regex fallback (when AST fails)
RE_DEF = re.compile(r'^\s*(async\s+def|def)\s+([A-Za-z_]\w*)\s*\(', re.M)
RE_CLASS = re.compile(r'^\s*class\s+([A-Za-z_]\w*)\s*[\(:]', re.M)
RE_IMPORT_LINE = re.compile(r'^\s*(import\s+[^\n]+|from\s+[^\n]+)', re.M)
RE_ADDARG = re.compile(r'\.add_argument\(\s*([\'"]-{1,2}[^\'"]+)[\'"]', re.M)
RE_TRIPLE_DOC = re.compile(r'^\s*[ruRU]{0,2}([\'"]{3})(.*?)\1', re.S | re.M)

# Heuristic header parsing (comment header)
RE_HEADER_KV = re.compile(
    r'^\s*#\s*(but|objectif|role|rôle|entrees|entrées|sorties|usage|note|notes)\s*:\s*(.+?)\s*$',
    re.I | re.M
)

# I/O heuristics
RE_FILE_EXT_IN_STR = re.compile(r'([A-Za-z0-9_\-\\/:\. ]+\.(json|csv|txt|log|png|jpg|jpeg|webp|pdf|odt|docx))', re.I)
RE_OPEN_CALL = re.compile(r'\bopen\s*\(', re.I)


# ---------------------------
# Data model
# ---------------------------

@dataclass
class PyFileInfo:
    rel_path: str
    abs_path: str
    size_bytes: int
    parse_mode: str                 # "ast" | "regex" | "syntax_error"
    has_main_guard: bool
    module_doc: str
    top_comment: str
    header_kv: Dict[str, str]
    imports: List[str]
    functions: List[str]
    classes: List[str]
    argparse_args: List[str]
    io_hints: List[str]
    guessed_role: str


# ---------------------------
# Utilities
# ---------------------------

def read_text_safe(p: Path) -> str:
    try:
        return p.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return p.read_text(encoding="cp1252", errors="replace")



def fix_mojibake(s: str) -> str:
    """Répare les accents quand un texte UTF-8 a été 'mojibaké' (ex: 'Ã©' au lieu de 'é')."""
    if not s:
        return s
    # Heuristique: présence de marqueurs fréquents
    if ("Ã" in s) or ("â€" in s) or ("Â" in s):
        try:
            repaired = s.encode("latin1").decode("utf-8")
            # Accepte si on réduit nettement les marqueurs
            if repaired and repaired.count("Ã") < s.count("Ã"):
                return repaired
        except Exception:
            pass
    return s


def extract_top_comment(src: str, max_lines: int = 60) -> str:
    """Extrait un 'résumé' en comment-header (sans prendre le shebang / coding)."""
    lines = src.splitlines()
    out: List[str] = []
    for line in lines[:max_lines]:
        s = line.rstrip()
        if not s.strip():
            if out:
                break
            continue

        ls = s.lstrip()
        if ls.startswith("#"):
            # ignore shebang + coding
            if ls.startswith("#!"):
                continue
            if ("coding" in ls) and ("-*- coding" in ls or "coding:" in ls):
                continue

            content = ls[1:].strip()
            # ignore separators like "-----"
            if content and set(content) <= {"-", "=", "_"}:
                continue

            out.append(content)
            continue
        break

    return fix_mojibake(" ".join(out).strip())




def extract_header_kv(src: str) -> Dict[str, str]:
    """Lit les clés d'en-tête du style '# But: ...', '# Usage: ...', etc."""
    out: Dict[str, str] = {}
    for m in RE_HEADER_KV.finditer(src[:6000]):  # only early section
        k = m.group(1).strip().lower()
        v = fix_mojibake(m.group(2).strip())
        # Normalize keys
        k = {
            "objectif": "but",
            "role": "rôle",
            "rôle": "rôle",
            "entrees": "entrées",
            "entrées": "entrées",
            "sorties": "sorties",
            "usage": "usage",
            "note": "notes",
            "notes": "notes",
            "but": "but",
        }.get(k, k)
        if k not in out:
            out[k] = v
    return out



def find_main_guard_ast(tree: ast.AST) -> bool:
    for node in ast.walk(tree):
        if isinstance(node, ast.If):
            try:
                test = ast.unparse(node.test)
            except Exception:
                continue
            if "__name__" in test and "__main__" in test:
                return True
    return False


def extract_imports_ast(tree: ast.AST) -> List[str]:
    imps: List[str] = []
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            for n in node.names:
                imps.append(n.name)
        elif isinstance(node, ast.ImportFrom):
            mod = node.module or ""
            for n in node.names:
                imps.append(f"{mod}.{n.name}" if mod else n.name)
    return sorted(set(imps))


def extract_defs_ast(tree: ast.AST) -> Tuple[List[str], List[str]]:
    funcs: List[str] = []
    clss: List[str] = []
    body = getattr(tree, "body", [])
    for node in body:
        if isinstance(node, ast.FunctionDef):
            funcs.append(node.name)
        elif isinstance(node, ast.AsyncFunctionDef):
            funcs.append(node.name + " (async)")
        elif isinstance(node, ast.ClassDef):
            clss.append(node.name)
    return funcs, clss



def extract_module_doc_ast(tree: ast.AST) -> str:
    d = ast.get_docstring(tree) or ""
    d = " ".join(d.split()).strip()
    return fix_mojibake(d)



def extract_argparse_args_ast(src: str, tree: ast.AST) -> List[str]:
    if "argparse" not in src:
        return []
    args: List[str] = []
    for node in ast.walk(tree):
        if isinstance(node, ast.Call):
            try:
                fn = ast.unparse(node.func)
            except Exception:
                continue
            if fn.endswith(".add_argument"):
                for a in node.args:
                    if isinstance(a, ast.Constant) and isinstance(a.value, str):
                        s = a.value.strip()
                        if s.startswith("-"):
                            args.append(s)
    return sorted(set(args))



def extract_module_doc_fallback(src: str) -> str:
    """Fallback docstring: ne prend que la docstring de module (en tout début de fichier)."""
    lines = src.splitlines()
    i = 0

    # skip shebang + encoding comments
    while i < len(lines):
        s = lines[i].lstrip()
        if i == 0 and s.startswith("#!"):
            i += 1
            continue
        if s.startswith("#") and ("coding:" in s or "-*- coding" in s):
            i += 1
            continue
        # skip blanks + comments
        if not s.strip() or s.startswith("#"):
            i += 1
            continue
        break

    rest = "\n".join(lines[i:])
    m = re.match(r'^\s*[rRuU]{0,2}([\'"]{3})(.*?)\1', rest, flags=re.S)
    if not m:
        return ""
    doc = " ".join(m.group(2).strip().split()).strip()
    return fix_mojibake(doc)


def fallback_parse_regex(src: str) -> Tuple[str, List[str], List[str], List[str], List[str], bool]:
    module_doc = extract_module_doc_fallback(src)

    functions = [m.group(2) for m in RE_DEF.finditer(src)]
    classes = [m.group(1) for m in RE_CLASS.finditer(src)]

    imports_raw = [m.group(1).strip() for m in RE_IMPORT_LINE.finditer(src)]
    imports = sorted(set(imports_raw))

    argparse_args = sorted(set(m.group(1).strip() for m in RE_ADDARG.finditer(src)))

    has_main_guard = ("__name__" in src) and ("__main__" in src)
    return module_doc, functions, classes, imports, argparse_args, has_main_guard



def guess_role(path: Path, rel_path: str, imports: List[str]) -> str:
    name = path.name.lower()
    rel = rel_path.lower()
    imp_join = " ".join(imports).lower()

    if "engine" in rel or name.startswith("ddm_p"):
        return "engine/core gameplay"
    if "server" in rel or "flask" in imp_join:
        return "server/api or web"
    if "tools" in rel and ("autorun" in rel or "runlab" in rel):
        return "automation/runs"
    if "policy" in rel or "policy" in name or "ai_policy" in name:
        return "ai/policy analysis"
    if "analy" in name or "report" in rel or "timeline" in name or "quality" in name:
        return "analysis/reporting"
    if "smc" in name:
        return "inventory/scanner"
    return "misc utility"


def should_exclude(path: Path, excludes: set[str]) -> bool:
    parts = {p.lower() for p in path.parts}
    return any(x.lower() in parts for x in excludes)


def detect_io_hints(src: str, imports: List[str]) -> List[str]:
    hints: List[str] = []

    # imports-based hints
    imp = " ".join(imports).lower()
    if "json" in imp:
        hints.append("manipule JSON")
    if "csv" in imp or "pandas" in imp:
        hints.append("produit/consomme CSV")
    if "matplotlib" in imp:
        hints.append("génère des graphiques")
    if "flask" in imp:
        hints.append("serveur web/API")
    if "duckdb" in imp:
        hints.append("utilise DuckDB")
    if "sqlite3" in imp:
        hints.append("utilise SQLite")

    # string-based file extension hints
    exts = set(m.group(2).lower() for m in RE_FILE_EXT_IN_STR.finditer(src[:20000]))
    # only keep meaningful ones
    for e in sorted(exts):
        if e in {"json", "csv", "log", "txt", "pdf", "docx", "odt"}:
            hints.append(f"référence des fichiers .{e}")

    # open() usage
    if RE_OPEN_CALL.search(src):
        hints.append("lit/écrit des fichiers (open)")

    # de-dup
    out: List[str] = []
    for h in hints:
        if h not in out:
            out.append(h)
    return out


# ---------------------------
# Scan
# ---------------------------

def scan(root: Path, targets: List[Path], excludes: set[str]) -> List[PyFileInfo]:
    out: List[PyFileInfo] = []
    for target in targets:
        if not target.exists():
            continue
        for p in target.rglob(f"*{PY_EXT}"):
            if should_exclude(p, excludes):
                continue

            src = read_text_safe(p)
            top_comment = extract_top_comment(src)
            header_kv = extract_header_kv(src)

            rel = str(p.relative_to(root))
            size = p.stat().st_size

            parse_mode = "ast"
            module_doc = ""
            imports: List[str] = []
            funcs: List[str] = []
            clss: List[str] = []
            ap_args: List[str] = []
            has_main = False

            try:
                tree = ast.parse(src)
                module_doc = extract_module_doc_ast(tree)
                imports = extract_imports_ast(tree)
                funcs, clss = extract_defs_ast(tree)
                ap_args = extract_argparse_args_ast(src, tree)
                has_main = find_main_guard_ast(tree)
            except SyntaxError:
                # fallback regex parsing
                parse_mode = "regex"
                module_doc, funcs, clss, imports, ap_args, has_main = fallback_parse_regex(src)
            except Exception:
                # unknown fatal parse issue
                parse_mode = "syntax_error"

            # IO hints
            io_hints = detect_io_hints(src, imports)

            role = guess_role(p, rel, imports)

            out.append(PyFileInfo(
                rel_path=rel,
                abs_path=str(p),
                size_bytes=size,
                parse_mode=parse_mode,
                has_main_guard=has_main,
                module_doc=module_doc,
                top_comment=top_comment,
                header_kv=header_kv,
                imports=imports,
                functions=funcs,
                classes=clss,
                argparse_args=ap_args,
                io_hints=io_hints,
                guessed_role=role,
            ))

    out.sort(key=lambda x: x.rel_path.lower())
    return out


# ---------------------------
# Output renderers
# ---------------------------


# ---------------------------
# CMD.txt integration (cheat-sheet)
# ---------------------------

def infer_cmd_file(root: Path) -> Optional[Path]:
    """Trouve un CMD.txt plausible (souvent dans ../Admin/CMD.txt)."""
    candidates = [
        root / "Admin" / "CMD.txt",
        root.parent / "Admin" / "CMD.txt",
        root / "CMD.txt",
        root.parent / "CMD.txt",
    ]
    for p in candidates:
        if p.exists():
            return p
    return None


def parse_cmd_blocks(text: str) -> List[List[str]]:
    """Découpe un CMD.txt en blocs (séparateurs / lignes vides)."""
    blocks: List[List[str]] = []
    cur: List[str] = []
    for raw in text.splitlines():
        line = raw.rstrip()
        stripped = line.strip()
        is_sep = stripped and set(stripped) <= {"-", "=", "_"}
        if (not stripped) or is_sep:
            if cur:
                blocks.append(cur)
                cur = []
            continue
        cur.append(line)
    if cur:
        blocks.append(cur)
    return blocks


def cmd_key_from_block(block: List[str]) -> Optional[str]:
    """Tente d'associer un bloc de commande à un script (clé = basename.py)."""
    joined = " ".join(block)
    joined = joined.replace("^", " ")
    joined = " ".join(joined.split())

    m = re.search(r"\bpython\s+-m\s+([A-Za-z0-9_\.]+)", joined)
    if m:
        mod = m.group(1).split(".")[-1]
        return f"{mod}.py"

    m = re.search(r"\bpython\s+['\"]?([^'\"\s]+\.py)['\"]?", joined)
    if m:
        return Path(m.group(1)).name

    return None


def load_cmd_map(root: Path, cmd_file: Optional[str] = None) -> Dict[str, List[str]]:
    """Retourne {script_basename.py: [bloc1, bloc2, ...]}"""
    p: Optional[Path] = None
    if cmd_file:
        p = Path(cmd_file)
        if not p.is_absolute():
            p = (root / p).resolve()
    else:
        p = infer_cmd_file(root)

    if not p or not p.exists():
        return {}

    txt = read_text_safe(p)
    blocks = parse_cmd_blocks(txt)

    out: Dict[str, List[str]] = {}
    for b in blocks:
        key = cmd_key_from_block(b)
        if not key:
            continue
        out.setdefault(key, []).append("\n".join(b).strip())
    return out


def short_desc(it: PyFileInfo, max_len: int = 140) -> str:
    d = it.header_kv.get("but") or it.module_doc or it.top_comment or ""
    d = fix_mojibake(d)
    if len(d) <= max_len:
        return d
    return d[:max_len - 1].rstrip() + "…"


def suggest_commands(it: PyFileInfo, root: Path, cmd_map: Dict[str, List[str]]) -> List[str]:
    """Renvoie une liste de commandes (blocs) pertinentes pour lancer le script."""
    # 1) CMD.txt (prioritaire)
    key = Path(it.abs_path).name
    if key in cmd_map:
        return cmd_map[key]

    # 2) Suggestion générique
    rel = it.rel_path.replace("/", "\\")
    if it.has_main_guard:
        return [f"cd /d {root}\npython {rel} --help"]
    return []

def to_markdown(items: List[PyFileInfo], root: Path) -> str:
    lines: List[str] = []
    lines.append(f"# CODEMAP — {root}")
    lines.append("")
    lines.append(f"Fichiers Python indexés : **{len(items)}**")
    lines.append("")

    # stats parse modes
    modes: Dict[str, int] = {}
    for it in items:
        modes[it.parse_mode] = modes.get(it.parse_mode, 0) + 1
    lines.append("## Qualité d’indexation")
    for k in sorted(modes):
        lines.append(f"- `{k}` : {modes[k]}")
    lines.append("")

    # group by role
    by_role: Dict[str, List[PyFileInfo]] = {}
    for it in items:
        by_role.setdefault(it.guessed_role, []).append(it)

    lines.append("## Sommaire par rôle")
    for role in sorted(by_role.keys()):
        lines.append(f"- **{role}** ({len(by_role[role])})")
    lines.append("")

    def short(s: str, n: int = 260) -> str:
        s = fix_mojibake(s or "").strip()
        return s if len(s) <= n else (s[:n].rstrip() + "…")

    for role in sorted(by_role.keys()):
        lines.append(f"## {role}")
        lines.append("")
        for it in by_role[role]:
            lines.append(f"### `{it.rel_path}`")
            lines.append(f"- Taille: `{it.size_bytes}` bytes | parse: `{it.parse_mode}` | main-guard: `{it.has_main_guard}`")

            # Prefer header "but", then docstring, then top comment
            if it.header_kv.get("but"):
                lines.append(f"- But: {short(it.header_kv['but'])}")
            elif it.module_doc:
                lines.append(f"- Doc: {short(it.module_doc)}")
            elif it.top_comment:
                lines.append(f"- Comment: {short(it.top_comment)}")

            if it.header_kv.get("usage"):
                lines.append(f"- Usage: `{it.header_kv['usage']}`")

            if it.argparse_args:
                lines.append(f"- CLI: {' '.join(sorted({a.strip() for a in it.argparse_args if a.strip()}))}")

            if it.io_hints:
                lines.append(f"- I/O: {', '.join(it.io_hints)}")

            if it.classes:
                lines.append(f"- Classes: {', '.join(it.classes[:30])}{'…' if len(it.classes) > 30 else ''}")

            if it.functions:
                lines.append(f"- Fonctions: {', '.join(it.functions[:40])}{'…' if len(it.functions) > 40 else ''}")

            # imports can be very verbose (especially regex mode)
            if it.imports:
                show = it.imports[:16]
                extra = "" if len(it.imports) <= 16 else f" (+{len(it.imports)-16})"
                lines.append(f"- Imports: {', '.join(show)}{extra}")

            # Header notes if present
            if it.header_kv.get("notes"):
                lines.append(f"- Notes: {short(it.header_kv['notes'])}")

            lines.append("")
    return "\n".join(lines)



def to_guide_markdown(items: List[PyFileInfo], root: Path, cmd_map: Dict[str, List[str]]) -> str:
    """Version plus 'guidée' et explicite, en FR."""
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    lines: List[str] = []
    lines.append("# Guide des scripts Proto_DDM (FR)")
    lines.append(f"_Généré le {now} — Racine scannée: `{root}` — Scripts indexés: **{len(items)}**_")
    lines.append("")
    lines.append("## Comment t'en servir")
    lines.append("- **Trouver un script** : va dans l'index par catégorie ci-dessous.")
    lines.append("- **Le lancer** : regarde la section **Exécution** (commande CMD.txt si dispo, sinon commande générique).")
    lines.append("- **Quand tu hésites** : lance `--help` (quand disponible).")
    lines.append("")
    lines.append("## Index par catégorie")

    groups: Dict[str, List[PyFileInfo]] = {}
    for it in items:
        role = simplify_role(it.guessed_role)
        groups.setdefault(role, []).append(it)

    for role in sorted(groups.keys()):
        lines.append(f"### {role} ({len(groups[role])})")
        for it in sorted(groups[role], key=lambda x: x.rel_path.lower()):
            lines.append(f"- **{it.rel_path}** — {short_desc(it)}")
        lines.append("")

    lines.append("## Détails script par script")
    for role in sorted(groups.keys()):
        lines.append(f"### {role}")
        for it in sorted(groups[role], key=lambda x: x.rel_path.lower()):
            lines.append(f"#### `{it.rel_path}`")
            lines.append(f"- Taille: `{it.size_bytes}` bytes | parse: `{it.parse_mode}` | main-guard: `{it.has_main_guard}`")
            d = it.header_kv.get("but") or it.module_doc or it.top_comment or ""
            d = fix_mojibake(d).strip()
            if d:
                lines.append(f"- **À quoi ça sert :** {d}")
            if it.argparse_args:
                lines.append(f"- **Options CLI détectées :** " + ", ".join(sorted({a.strip() for a in it.argparse_args if a.strip()})))
            if it.io_hints:
                lines.append(f"- **Entrées / Sorties (indices) :** " + ", ".join(it.io_hints))
            # show command blocks
            cmds = suggest_commands(it, root, cmd_map)
            if cmds:
                lines.append("- **Exécution :**")
                for block in cmds[:3]:
                    lines.append("```bat")
                    lines.append(block)
                    lines.append("```")
            else:
                lines.append("- **Exécution :** (pas de commande détectée — probablement un module utilitaire)")
            # show main functions/classes (limit)
            if it.functions:
                fn = ", ".join(it.functions[:12])
                if len(it.functions) > 12:
                    fn += ", …"
                lines.append(f"- **Fonctions repérées :** {fn}")
            if it.classes:
                cl = ", ".join(it.classes[:10])
                if len(it.classes) > 10:
                    cl += ", …"
                lines.append(f"- **Classes repérées :** {cl}")
            lines.append("")
        lines.append("")

    return "\n".join(lines) + "\n"

def simplify_role(role: str) -> str:
    return {
        "engine/core gameplay": "Moteur du jeu (gameplay)",
        "automation/runs": "Lancement automatique de parties",
        "analysis/reporting": "Analyse / rapports",
        "ai/policy analysis": "IA / stratégies (policies)",
        "server/api or web": "Serveur web / interface",
        "inventory/scanner": "Inventaire / scan de fichiers",
        "misc utility": "Utilitaires divers",
    }.get(role, role)


def to_txt_readable(items: List[PyFileInfo], root: Path) -> str:
    out: List[str] = []
    out.append("CODEMAP — version lecture (plus simple)\n")
    out.append(f"Projet: {root}")
    out.append(f"Scripts indexés: {len(items)}\n")

    by_role: Dict[str, List[PyFileInfo]] = {}
    for it in items:
        by_role.setdefault(it.guessed_role, []).append(it)

    for role in sorted(by_role.keys()):
        out.append(f"\n== {simplify_role(role)} ({len(by_role[role])}) ==\n")
        for it in by_role[role]:
            out.append(f"- {it.rel_path}")

            desc = fix_mojibake(it.header_kv.get("but") or it.module_doc or it.top_comment or "")
            if desc:
                out.append(f"  À quoi ça sert : {desc[:220] + ('…' if len(desc) > 220 else '')}")

            if it.header_kv.get("usage"):
                out.append(f"  Comment l’utiliser : {it.header_kv['usage']}")

            if it.argparse_args:
                out.append(f"  Options (CLI) : {' '.join(it.argparse_args)}")

            if it.io_hints:
                out.append(f"  Indices I/O : {', '.join(it.io_hints)}")

            if it.has_main_guard:
                out.append("  Exécutable directement : oui")

            out.append("")
    return "\n".join(out)


def write_docx(text: str, out_path: Path) -> None:
    # optional dependency (installed in your environment typically)
    try:
        from docx import Document
    except Exception as e:
        raise RuntimeError("python-docx n'est pas disponible. Installe-le ou enlève 'docx' des formats.") from e

    doc = Document()
    for line in text.splitlines():
        if line.startswith("== ") and line.endswith(" =="):
            doc.add_heading(line.strip("= ").strip(), level=1)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.strip() == "":
            doc.add_paragraph("")
        else:
            doc.add_paragraph(line)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


# ---------------------------
# Interactive menu
# ---------------------------

def ask(prompt: str, default: Optional[str] = None) -> str:
    if default is not None:
        s = input(f"{prompt} [{default}] > ").strip()
        return s if s else default
    return input(f"{prompt} > ").strip()


def ask_yes_no(prompt: str, default: bool = True) -> bool:
    d = "Y/n" if default else "y/N"
    s = input(f"{prompt} ({d}) > ").strip().lower()
    if not s:
        return default
    return s in {"y", "yes", "o", "oui"}


def ask_choice(prompt: str, choices: List[str], default_idx: int = 0) -> str:
    print(prompt)
    for i, c in enumerate(choices, 1):
        mark = " (défaut)" if i - 1 == default_idx else ""
        print(f"  {i}. {c}{mark}")
    while True:
        s = input("> ").strip()
        if not s:
            return choices[default_idx]
        if s.isdigit():
            idx = int(s) - 1
            if 0 <= idx < len(choices):
                return choices[idx]
        print("Choix invalide. Réessaie.")



def parse_formats(s: str) -> List[str]:
    raw = [x.strip().lower() for x in s.split(",") if x.strip()]
    ok: List[str] = []
    allowed = {"md", "json", "txt", "docx", "guide"}
    for x in raw:
        if x in allowed and x not in ok:
            ok.append(x)
    return ok




def interactive_config(script_path: Path) -> Tuple[Path, List[str], str, List[str], set[str]]:
    default_root = script_path.resolve().parents[2]  # ...\tools\scripts -> ...\Proto_DDM
    root_str = ask("Racine du projet", str(default_root))
    root = Path(root_str).resolve()

    target_mode = ask_choice(
        "Qu’est-ce que tu veux indexer ?",
        [
            "Projet entier (recommandé)",
            "engine seulement",
            "tools seulement",
            "server seulement",
            "Sélection manuelle (chemin relatif)",
        ],
        default_idx=0
    )

    targets: List[str] = []
    if target_mode.startswith("Projet entier"):
        targets = ["."]
    elif target_mode.startswith("engine"):
        targets = ["engine"]
    elif target_mode.startswith("tools"):
        targets = ["tools"]
    elif target_mode.startswith("server"):
        targets = ["server"]
    else:
        rel = ask("Chemin relatif à indexer (ex: tools\\scripts)", "tools")
        targets = [rel]

    out_dir = ask("Dossier de sortie (relatif au root)", DEFAULT_OUT_DIR)

    formats_str = ask("Formats (md,json,txt,docx) séparés par virgules", ",".join(DEFAULT_FORMATS))
    formats = parse_formats(formats_str) or DEFAULT_FORMATS

    # Excludes
    excludes = set(DEFAULT_EXCLUDES)
    if ask_yes_no("Afficher et modifier les exclusions ?", default=False):
        print("Exclusions actuelles:", ", ".join(sorted(excludes)))
        while True:
            op = ask_choice("Action exclusions", ["Ajouter", "Retirer", "OK"], default_idx=2)
            if op == "OK":
                break
            if op == "Ajouter":
                ex = ask("Nom de dossier à exclure", "")
                if ex:
                    excludes.add(ex)
            elif op == "Retirer":
                ex = ask("Nom de dossier à retirer des exclusions", "")
                if ex and ex in excludes:
                    excludes.remove(ex)
        print("Exclusions finales:", ", ".join(sorted(excludes)))

    return root, targets, out_dir, formats, excludes


# ---------------------------
# Main
# ---------------------------

def build_targets(root: Path, targets: List[str]) -> List[Path]:
    out: List[Path] = []
    for t in targets:
        t = t.strip()
        if t in {".", ""}:
            out.append(root)
        else:
            out.append((root / t).resolve())
    # de-dup
    dedup: List[Path] = []
    seen = set()
    for p in out:
        if str(p).lower() not in seen:
            dedup.append(p)
            seen.add(str(p).lower())
    return dedup


def main() -> int:
    ap = argparse.ArgumentParser(description="Indexe les .py et génère CODEMAP (md/json/txt/docx/guide) + menu interactif.")
    ap.add_argument("--interactive", action="store_true", help="Lance un menu interactif.")
    ap.add_argument("--root", help=r"Racine projet (ex: .)")
    ap.add_argument("--targets", default=".", help=r"Cibles relatives séparées par ; (ex: engine;tools\scripts). Défaut: '.'")
    ap.add_argument("--out-dir", default=DEFAULT_OUT_DIR, help="Dossier de sortie relatif à root. Défaut: docs")
    ap.add_argument("--formats", default="md,json", help="Formats: md,json,txt,docx,guide (séparés par virgules).")
    ap.add_argument("--cmd-file", default="", help=r"Chemin vers CMD.txt (optionnel). Si vide: auto (../Admin/CMD.txt).")
    ap.add_argument("--exclude", action="append", default=[], help="Dossier à exclure (répétable).")

    args = ap.parse_args()

    script_path = Path(__file__)

    if args.interactive or (not args.root):
        root, targets_rel, out_dir, formats, excludes = interactive_config(script_path)
    else:
        root = Path(args.root).resolve()
        targets_rel = [x.strip() for x in args.targets.split(";") if x.strip()]
        out_dir = args.out_dir
        formats = parse_formats(args.formats) or DEFAULT_FORMATS
        excludes = set(DEFAULT_EXCLUDES) | {x.strip() for x in args.exclude if x.strip()}

    if not root.exists():
        print("ERREUR: root introuvable:", root)
        return 2

    targets = build_targets(root, targets_rel)
    items = scan(root, targets, excludes)

    out_base = (root / out_dir).resolve()
    out_base.mkdir(parents=True, exist_ok=True)

    # Always JSON when requested
    if "json" in formats:
        (out_base / "CODEMAP.json").write_text(
            json.dumps([asdict(x) for x in items], indent=2, ensure_ascii=False),
            encoding="utf-8"
        )

    if "md" in formats:
        (out_base / "CODEMAP.md").write_text(
            to_markdown(items, root),
            encoding="utf-8"
        )

    if "txt" in formats or "docx" in formats:
        txt = to_txt_readable(items, root)
        if "txt" in formats:
            (out_base / "CODEMAP.txt").write_text(txt, encoding="utf-8")
        if "docx" in formats:
            write_docx(txt, out_base / "CODEMAP.docx")


    # Guide explicite (FR)
    if "guide" in formats:
        cmd_file = (args.cmd_file or "").strip() or None
        cmd_map = load_cmd_map(root, cmd_file=cmd_file)
        guide = to_guide_markdown(items, root, cmd_map)
        (out_base / "SCRIPTS_GUIDE_FR.md").write_text(guide, encoding="utf-8")

    # Summary
    modes: Dict[str, int] = {}
    for it in items:
        modes[it.parse_mode] = modes.get(it.parse_mode, 0) + 1

    print("OK — CODEMAP généré")
    print("Root:", root)
    print("Targets:", ", ".join(str(t) for t in targets))
    print("Out:", out_base)
    print("Formats:", ", ".join(formats))
    print("Parse modes:", ", ".join(f"{k}={v}" for k, v in sorted(modes.items())))

    return 0


if __name__ == "__main__":
    raise SystemExit(main())